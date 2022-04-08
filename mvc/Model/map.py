import csv
import math

import numpy as np
import openpyxl
import pandas as pd
from matplotlib import pyplot as plt

from mvc.Model.roof_profile import RoofProfile, RoofPoint
from mvc.Model.shape import Shape
from mvc.Model.size import Size
from mvc.Model.split import Split
from mvc.Model.surface import get_square_surface
from utils.filedialog import dict_from_json
from utils.geometry.point_in_polygon import check_polygon_in_polygon
from utils.json_in_out import JsonInOut
from utils.observer import Subject
from utils.transform_data_to_export import dict_update, transform_data

PointType = (float, float, int)


def pop_from_dict(dict_value: dict, name: str):
    dict_value[name] = False
    dict_value.pop(name)


def pop_from_dict_many(dict_value: dict, names: [str]):
    for name in names:
        pop_from_dict(dict_value, name)


class Map(Subject, JsonInOut):
    __slots__ = 'size', 'shapes', 'roof_profile', 'data', 'draw_speed', 'splits'

    def __init__(self, data: dict = None):
        super().__init__()
        self.__call__(data)

    def __call__(self, data: dict = None, *args, **kwargs):
        self.draw_speed = 'Fast'
        self.data = {}
        self.size = Size()
        self.roof_profile = RoofProfile()
        self.splits: [Split] = [Split(), Split()]
        self.shapes: [Shape] = list()

        if data:
            self.__load_from_dict(data)

    def add_layer(self, figure: Shape = None) -> Shape:

        if not figure:
            figure = Shape(size=self.size)
        if not figure._observers:
            figure._observers = self._observers

        self.shapes.append(figure)
        self.notify()
        return figure

    def delete_layer(self, index: int = None, figure: Shape = None):
        if index:
            self.shapes.pop(index)
        if figure:
            self.shapes.remove(figure)
        self.notify()

    def get_shapes(self) -> [Shape]:
        return self.shapes

    def get_visible_shapes(self) -> [Shape]:
        shapes_with_split = []
        for shape in [i for i in self.shapes if i.visible]:
            shapes_with_split = shapes_with_split + shape.splitting_shape(self.splits)
        return sorted(filter(lambda i: i.visible is True, shapes_with_split), key=lambda i: i.priority).__reversed__()

    def get_shape_with_part(self) -> [Shape]:
        layers = []
        for layer in self.get_shapes():
            layers.append(layer)
            layers = layers + [layer.parts_property[part_name] for part_name in layer.parts_property]
        return layers

    def get_as_dict(self) -> dict:
        map_dict = super(Map, self).get_as_dict()
        pop_from_dict_many(map_dict, ['data', 'draw_speed'])
        for d in map_dict:
            if d == 'shapes':
                for shape in map_dict[d]:
                    pop_from_dict_many(shape, ['size', 'split_shapes'])

                    for name in shape.get('parts_property'):
                        pop_names = ['size', 'layers', 'splits', 'split_shapes', 'parts_property']
                        pop_from_dict_many(shape['parts_property'][name], pop_names)
                    pop_layer_number = []

                    for lay, i in zip(shape['layers'], range(len(shape['layers']))):
                        if lay.get('primary') is True:
                            pop_names = ['size', 'pre_x', 'pre_y', 'start_x', 'start_y', 'splits', 'current_split']
                            pop_from_dict_many(lay, pop_names)
                        elif lay.get('primary') is False:
                            pop_layer_number.append(i)
                    for i in pop_layer_number.__reversed__():
                        shape['layers'].pop(i)
                    return map_dict

    def load_from_dict(self, load_dict: dict):
        self(load_dict)

    def __load_from_dict(self, load_dict: dict):
        for name_property in load_dict:
            if name_property == 'shapes':
                for lay in load_dict['shapes']:
                    self.add_layer(Shape(size=self.size, load_dict=lay))
            elif name_property == 'splits':
                self.splits = [Split(split) for split in load_dict['splits']]
            elif name_property == 'roof_profile':
                self.roof_profile.load_from_dict(load_dict[name_property])
            elif hasattr(self, name_property):
                if hasattr(getattr(self, name_property), 'load_from_dict'):
                    getattr(self, name_property).load_from_dict(load_dict[name_property])
                else:
                    self.__setattr__(name_property, load_dict[name_property])

    def load_from_json(self, path: str):
        map_dict = dict_from_json(path)
        self.load_from_dict(map_dict)

    def update_size(self):
        x = [a for b in [surf.x for shape in self.shapes for surf in shape.layers] for a in b] + \
            [self.size.x_constraints.start, self.size.x_constraints.end]
        y = [a for b in [surf.y for shape in self.shapes for surf in shape.layers] for a in b] + \
            [self.size.y_constraints.start, self.size.y_constraints.end]

        self.size.x_constraints.start, self.size.x_constraints.end = min(x), max(x)
        self.size.y_constraints.start, self.size.y_constraints.end = min(y), max(y)
        self.size.z_constraints.end = max([i.height for i in self.shapes] + [self.size.z_constraints.end])

    @property
    def height(self) -> int:
        if len(self.shapes) > 0:
            return max(self.shapes, key=lambda i: i.height).height
        else:
            return 0

    @property
    def height_with_offset(self) -> int:
        max_offset = int(max(self.roof_profile.points + [RoofPoint(x=0, z=0, y=0)], key=lambda i: i.z).z + 1)
        return max(self.shapes, key=lambda i: i.height_with_offset).height_with_offset + max_offset


class ExportMap:
    def __init__(self, map: Map):
        self.map = map
        self.repeat, self.map.data = {}, {}

    def __call__(self, *args, **kwargs) -> dict:
        for shape in self.map.shapes:
            shape.calc_intermediate_layers()
        return self.export()

    def export(self) -> dict:
        self.__init__(self.map)
        for shape in self.map.get_visible_shapes():
            data = self.calc_polygon_in_draw(shape)
            self.map.data[f'{shape.name}|{shape.sub_name}'] = \
                dict_update(self.map.data.get(shape.name), transform_data(data))

        self.map.data = {k: v for k, v in self.map.data.items() if v}
        return self.map.data

    def calc_polygon_in_draw(self, fig: Shape) -> []:
        roof = self.map.roof_profile.get_x_y_offset(base=max(self.map.size.x, self.map.size.y))
        roof = [[0 if fig.filler else j for j in i] for i in roof]
        x_size, y_size, z_size = self.map.size.x, self.map.size.y, int(fig.height + max(a for b in roof for a in b) + 1)
        data = np.zeros([x_size, y_size, z_size], dtype=bool)

        for lay in fig.layers:
            (x, y), lay_z = lay.scalable_curve, lay.z
            for x1 in range(fig.size.x):
                x1_c = math.ceil(x1)
                xx, rpo_x = [x1_c, x1_c + 1, x1_c, x1_c + 1], roof[x1]
                for y1 in range(fig.size.y):
                    y1_c = math.ceil(y1)
                    z1_offset, yy = int(lay_z + rpo_x[y1]), [y1_c, y1_c, y1_c + 1, y1_c + 1]
                    if check_polygon_in_polygon(x, y, xx, yy):
                        rep_name = f'{x1}-{y1}-{z1_offset}'
                        if self.repeat.get(rep_name) is None and 0 <= z1_offset < fig.size.z:
                            self.repeat[rep_name] = True
                            data[x1, y1, z1_offset] = True

        return self.correction_strong_mixing(data, fig.size.x, fig.size.y)

    def correction_strong_mixing(self, data: [], size_x, size_y) -> []:
        for x1, y1 in [(x1, y1) for x1 in range(size_x) for y1 in range(size_y)]:
            data_column, convert_val = data[x1][y1], []
            for i in range(len(data_column)):
                if not data_column[i]:
                    convert_val.append(i)
                if data_column[i]:
                    if len(convert_val) < 6:
                        for j in convert_val:
                            self.repeat[f'{x1}-{y1}-{j}'], data_column[j] = True, True
                    convert_val = []
        return data


def plot_roof(xs: [float], ys: [float], zs: [float]):
    plt.figure().add_subplot(111, projection='3d').scatter(xs, ys, zs)
    plt.show()


def fit_to_grid(data: {PointType}, x_size: int = 25, y_size: int = 25) -> {PointType}:
    fit_data = {}
    for i, j in [(i, j) for i in range(x_size) for j in range(y_size)]:
        if data.get(f'{i}-{j}') is None:
            for i1, j1 in [(i - 1, j - 1), (i + 1, j - 1), (i - 1, j + 1), (i + 1, j + 1)]:
                if data.get(f'{i1}-{i1}') is not None:
                    fit_data[f'{i}-{j}'] = data[f'{i1}-{i1}']
                    break
        else:
            fit_data[f'{i}-{j}'] = data[f'{i}-{j}']

    return fit_data


class ExportRoof(ExportMap):
    def __init__(self, data_map: Map, initial_depth=2000, step_depth=0.2, path: str = None):
        super(ExportRoof, self).__init__(data_map)
        self.initial_depth = initial_depth
        self.step_depth = step_depth
        self.template = {
            'id': 0,
            'cSurface': None,
            'seisVal': None,
            'iSurf': None,
            'i_index': None,
            'j_index': None,
            'cType': 'inMeters',
            'xCoord': '',
            'yCoord': ''
        }
        if path is None:
            path = 'test'
        path += '.csv'
        print(path)
        self.export_to_csv(path)

    def export_to_csv(self, path: str):
        main_layers = [s for s in self.map.shapes if s.filler is False]

        max_high_layers = max(main_layers, key=lambda i: max(i.layers, key=lambda i: i.z).z).layers
        max_z = max(max_high_layers, key=lambda i: i.z).z

        min_high_layers = min(main_layers, key=lambda i: min(i.layers, key=lambda i: i.z).z).layers
        min_z = min(min_high_layers, key=lambda i: i.z).z

        main_shapes, parts_property = [s for s in self.map.shapes if not s.filler], {}
        if len(main_shapes):
            parts_property = {k: Shape(self.map.size, load_dict=v.get_as_dict())
                              for k, v in main_shapes[0].parts_property.items()}

        shape = Shape(self.map.size)
        shape.parts_property = parts_property
        shape.add_layer(get_square_surface(self.map.size, max_z, 24.99))
        shapes = shape.splitting_shape(self.map.splits)

        points = {}
        for shape in shapes:
            roof = transform_data(self.calc_polygon_in_draw(shape))
            points.update({f'{k}-{k1}': (k, k1, v1[0]['s'])
                           for k, v in roof.items() for k1, v1 in v.items()})

        fit_to_grid(points, self.map.size.x, self.map.size.y)
        points = sorted(points.values(), key=lambda i: i[0])
        points = sorted(points, key=lambda i: i[1])
        df = self.data_prepare_for_export(points, max_z - min_z)
        df.to_csv(path, index=False)
        plot_roof(xs=df['i_index'], ys=df['j_index'], zs=df['seisVal'])

    def data_prepare_for_export(self, data: [PointType], thickness_of_the_formation: float) -> pd.DataFrame:
        pf = []
        template = self.template.copy()
        for x1, y1, z1 in data:
            template.update({
                'cSurface': 'Layer1',
                'seisVal': self.initial_depth + z1 * self.step_depth,
                'iSurf': 1,
                'i_index': x1 + 1,
                'j_index': y1 + 1,
            })
            pf.append(template.copy())
            template.update({
                'cSurface': 'Layer2',
                'seisVal': template['seisVal'] + thickness_of_the_formation * self.step_depth,
                'iSurf': 2,
            })
            pf.append(template.copy())

        return pd.DataFrame(pf, columns=template.keys())
